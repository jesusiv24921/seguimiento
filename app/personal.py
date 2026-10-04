"""Finanzas personales: reglas y persistencia Excel, independiente de Dash.

Importes en pesos enteros; movimientos y saldo se validan como una unidad.
El archivo separado evita competir con las escrituras de seguimiento.xlsx.
"""
from __future__ import annotations

from calendar import monthrange
from copy import deepcopy
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation
import hashlib
import json
import os
import time
from pathlib import Path
from threading import RLock
import uuid

import openpyxl

SCHEMAS = {
    "categorias": "id nombre tipo estado padre".split(),
    "ingresos": "id fecha descripcion categoria valor_planeado valor_real estado observaciones".split(),
    "gastos": "id fecha descripcion categoria valor_planeado valor_real estado tipo observaciones".split(),
    "deudas": "id nombre entidad saldo_inicial saldo_actual tasa_EA numero_cuotas cuotas_restantes fecha_inicio pago_programado estado".split(),
    "movimientos": "id fecha tipo_movimiento valor deuda capital observaciones".split(),
    "proyecciones": "id fecha disponible_proyectado abono_extraordinario observaciones".split(),
}
STATES = {
    "categorias": ["Activa", "Inactiva"],
    "ingresos": ["Planeado", "Recibido", "Parcial"],
    "gastos": ["Planeado", "Pagado", "Parcial", "Pendiente"],
    "deudas": ["Activa", "Pagada", "Suspendida"],
}
MOVEMENTS = ["Cuota normal", "Abono dirigido a capital", "Pago total", "Interés", "Otro"]
MONEY = {"valor_planeado", "valor_real", "saldo_inicial", "saldo_actual", "pago_programado",
         "valor", "capital", "disponible_proyectado", "abono_extraordinario"}
_LOCK = RLock()


def today():
    return datetime.now(timezone(timedelta(hours=-5))).date()


def storage_path():
    source = os.environ.get("SEGUIMIENTO_EXCEL_PATH")
    parent = Path(source).parent if source else Path(__file__).resolve().parent.parent
    return parent / "personal.xlsx"


def empty():
    return {key: [] for key in SCHEMAS}


def revision(data):
    return hashlib.sha256(json.dumps(data, sort_keys=True, ensure_ascii=False).encode()).hexdigest()


def load(path=None):
    path = Path(path) if path else storage_path()
    if not path.exists():
        return empty()
    result = empty()
    with _LOCK:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
        try:
            for key, columns in SCHEMAS.items():
                if key not in wb.sheetnames:
                    if key == "categorias":
                        continue  # Versión anterior: categorías escritas como texto.
                    raise ValueError(f"Falta la hoja {key} en personal.xlsx.")
                rows = wb[key].iter_rows(values_only=True)
                if list(next(rows, ())) != columns:
                    raise ValueError(f"Las columnas de {key} no corresponden a la versión actual.")
                result[key] = [dict(zip(columns, row)) for row in rows if any(v is not None for v in row)]
                for record in result[key]:
                    for column in columns:
                        if record[column] is None:
                            record[column] = 0 if column in MONEY or column in {"tasa_EA", "numero_cuotas", "cuotas_restantes"} else ""
        finally:
            wb.close()
        if "categorias" not in wb.sheetnames:
            migrate_categories(result)
    return result


def migrate_categories(data):
    """Migración en memoria, determinista; se persiste en el siguiente guardado."""
    for kind, category_type in (("ingresos", "Ingreso"), ("gastos", "Gasto")):
        for row in data[kind]:
            name = " ".join(str(row["categoria"]).split())
            if not name:
                raise ValueError("Hay registros sin categoría en el archivo anterior. Completa su categoría antes de migrar.")
            category_id = str(uuid.uuid5(uuid.NAMESPACE_URL, "seguimiento/personal/" + category_type + "/" + name.casefold()))
            if not any(c["id"] == category_id for c in data["categorias"]):
                data["categorias"].append(dict(id=category_id, nombre=name, tipo=category_type, estado="Activa", padre=""))
            row["categoria"] = category_id


def category_label(data, category_id):
    catalog = {c["id"]: c for c in data["categorias"]}
    names, seen = [], set()
    while category_id and category_id in catalog and category_id not in seen:
        seen.add(category_id)
        category = catalog[category_id]
        names.insert(0, category["nombre"])
        category_id = category["padre"]
    return " / ".join(names) or "Categoría no encontrada"


def category_active(data, category_id):
    catalog = {c["id"]: c for c in data["categorias"]}
    seen = set()
    while category_id:
        if category_id not in catalog or category_id in seen:
            return False
        seen.add(category_id)
        category = catalog[category_id]
        if category["estado"] != "Activa":
            return False
        category_id = category["padre"]
    return True


def category_options(data, kind, current=None):
    category_type = {"ingresos": "Ingreso", "gastos": "Gasto"}[kind]
    return [{"label": category_label(data, c["id"]) + (" (inactiva)" if not category_active(data, c["id"]) else ""),
             "value": c["id"]} for c in sorted(data["categorias"], key=lambda c: category_label(data, c["id"]).casefold())
            if c["tipo"] == category_type and (category_active(data, c["id"]) or c["id"] == current)]


def validate_categories(data):
    catalog = {c["id"]: c for c in data["categorias"]}
    identities = set()
    for c in catalog.values():
        identity = (c["tipo"], c["padre"], " ".join(c["nombre"].split()).casefold())
        if identity in identities:
            raise ValueError("Ya existe una categoría con ese nombre, tipo y categoría principal.")
        identities.add(identity)
        parent = c["padre"]
        seen = {c["id"]}
        while parent:
            if parent in seen:
                raise ValueError("Las subcategorías no pueden formar ciclos.")
            seen.add(parent)
            if parent not in catalog or catalog[parent]["tipo"] != c["tipo"]:
                raise ValueError("La categoría principal debe existir y ser del mismo tipo.")
            parent = catalog[parent]["padre"]
    for kind, category_type in (("ingresos", "Ingreso"), ("gastos", "Gasto")):
        for row in data[kind]:
            category = catalog.get(row["categoria"])
            if category is None or category["tipo"] != category_type:
                raise ValueError("Cada registro debe tener una categoría existente de su tipo. No cambies el tipo de una categoría con histórico.")


def amount(value, label="Valor"):
    try:
        number = Decimal(str(0 if value in (None, "") else value))
    except InvalidOperation:
        raise ValueError(f"{label}: ingresa un número válido.") from None
    if not number.is_finite() or number < 0 or number != number.to_integral_value() or number > 10**15:
        raise ValueError(f"{label}: usa pesos enteros entre 0 y 1.000.000.000.000.000.")
    return int(number)


def valid_date(value):
    try:
        return date.fromisoformat(str(value)).isoformat()
    except (ValueError, TypeError):
        raise ValueError("Ingresa una fecha válida (AAAA-MM-DD).") from None


def normalize(kind, row):
    out = {key: row.get(key) for key in SCHEMAS[kind]}
    out["id"] = str(out["id"] or uuid.uuid4())
    for key in out:
        if key in MONEY:
            out[key] = amount(out[key], key)
        elif key not in {"numero_cuotas", "cuotas_restantes", "tasa_EA"}:
            out[key] = str(out[key] or "").strip()
    if kind != "categorias":
        date_key = "fecha_inicio" if kind == "deudas" else "fecha"
        out[date_key] = valid_date(out[date_key])
    if kind in STATES and out["estado"] not in STATES[kind]:
        raise ValueError("Selecciona un estado válido.")
    required = {"categorias": ["nombre", "tipo"], "ingresos": ["descripcion", "categoria"], "gastos": ["descripcion", "categoria"],
                "deudas": ["nombre", "entidad"], "movimientos": ["deuda"], "proyecciones": []}[kind]
    if any(not out[key] for key in required):
        raise ValueError("Completa los campos obligatorios: " + ", ".join(required))
    if kind == "categorias":
        out["nombre"] = " ".join(out["nombre"].split())
        if out["tipo"] not in ("Ingreso", "Gasto"):
            raise ValueError("Selecciona Ingreso o Gasto para la categoría.")
    if kind == "gastos" and out["tipo"] not in ["Fijo", "Variable", "Deuda"]:
        raise ValueError("Selecciona un tipo de gasto válido.")
    if kind == "deudas":
        for key in ("numero_cuotas", "cuotas_restantes"):
            out[key] = amount(out[key], key)
        if out["cuotas_restantes"] > out["numero_cuotas"]:
            raise ValueError("Las cuotas restantes no pueden superar las cuotas totales.")
        try:
            rate = Decimal(str(out["tasa_EA"] or 0))
        except InvalidOperation:
            raise ValueError("La tasa E.A. no es válida.") from None
        if not rate.is_finite() or rate < 0 or rate > 1000:
            raise ValueError("La tasa E.A. debe estar entre 0 y 1000%.")
        out["tasa_EA"] = float(rate)
    if kind == "movimientos":
        if out["tipo_movimiento"] not in MOVEMENTS or out["valor"] <= 0:
            raise ValueError("Selecciona el movimiento e ingresa un valor mayor a cero.")
        if out["tipo_movimiento"] in ["Abono dirigido a capital", "Pago total"]:
            out["capital"] = out["valor"]
        elif out["tipo_movimiento"] == "Interés":
            out["capital"] = 0
        if out["capital"] > out["valor"]:
            raise ValueError("El capital no puede superar el valor pagado.")
    if kind == "proyecciones":
        out["fecha"] = out["fecha"][:7] + "-01"
    return out


def debt_balances(data, cutoff="9999-12-31"):
    balances = {d["id"]: d["saldo_inicial"] for d in data["deudas"] if d["fecha_inicio"] <= cutoff}
    debts = {d["id"]: d for d in data["deudas"]}
    for m in sorted(data["movimientos"], key=lambda r: r["fecha"]):
        if m["deuda"] not in debts:
            raise ValueError("El movimiento debe asociarse a una deuda existente.")
        if m["fecha"] < debts[m["deuda"]]["fecha_inicio"]:
            raise ValueError("Un movimiento no puede ser anterior al inicio de la deuda.")
        if m["fecha"] > cutoff:
            continue
        balance = balances[m["deuda"]]
        if m["tipo_movimiento"] == "Interés":
            balance += m["valor"]
        else:
            if m["tipo_movimiento"] == "Pago total" and m["capital"] != balance:
                raise ValueError("El pago total debe coincidir con el saldo en esa fecha.")
            balance -= m["capital"]
        if balance < 0:
            raise ValueError("El abono supera el saldo de la deuda; el saldo no puede ser negativo.")
        balances[m["deuda"]] = balance
    return balances


def mutate(kind, row=None, delete_id=None, expected=None, path=None):
    """Transacción dentro del único worker de Render, con revisión optimista."""
    path = Path(path) if path else storage_path()
    with _LOCK:
        current = load(path)
        if expected is not None and expected != revision(current):
            raise ValueError("Los datos cambiaron en otra pestaña. Recarga antes de guardar.")
        data = deepcopy(current)
        if delete_id:
            if not any(r["id"] == delete_id for r in data[kind]):
                raise ValueError("El registro ya no existe.")
            if kind == "deudas" and any(m["deuda"] == delete_id for m in data["movimientos"]):
                raise ValueError("Elimina primero los movimientos asociados a esta deuda.")
            if kind == "categorias":
                if any(r["categoria"] == delete_id for k in ("ingresos", "gastos") for r in data[k]):
                    raise ValueError("La categoría tiene histórico y no puede eliminarse. Cambia su estado a Inactiva.")
                if any(c["padre"] == delete_id for c in data["categorias"]):
                    raise ValueError("La categoría tiene subcategorías. Puedes desactivarla para conservar la estructura.")
            data[kind] = [r for r in data[kind] if r["id"] != delete_id]
        else:
            normalized = normalize(kind, row)
            if kind in ("ingresos", "gastos"):
                previous = next((r for r in current[kind] if r["id"] == normalized["id"]), None)
                if not category_active(data, normalized["categoria"]) and (not previous or previous["categoria"] != normalized["categoria"]):
                    raise ValueError("Selecciona una categoría activa. Las inactivas solo se conservan en registros ya asociados.")
            if row.get("id") and not any(r["id"] == row["id"] for r in data[kind]):
                raise ValueError("El registro ya no existe. Recarga los datos.")
            others = [r for r in data[kind] if r["id"] != normalized["id"]]
            comparable = lambda r: {k: v for k, v in r.items() if k not in {"id", "saldo_actual"}}
            if any(comparable(r) == comparable(normalized) for r in others):
                raise ValueError("Ya existe un registro idéntico; no se guardó el duplicado.")
            if kind == "proyecciones" and any(r["fecha"] == normalized["fecha"] for r in others):
                raise ValueError("Ya existe una proyección para ese mes. Edítala.")
            data[kind] = [normalized if r["id"] == normalized["id"] else r for r in data[kind]] if row.get("id") else others + [normalized]
        validate_categories(data)
        balances = debt_balances(data)
        for d in data["deudas"]:
            d["saldo_actual"] = balances[d["id"]]
            if d["estado"] == "Pagada" and d["saldo_actual"]:
                raise ValueError("Una deuda con saldo pendiente no puede marcarse como pagada.")
        if kind == "proyecciones" and not delete_id:
            project(data)
        path.parent.mkdir(parents=True, exist_ok=True)
        temp = path.with_name(f".{path.stem}-{uuid.uuid4().hex}.xlsx")
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        try:
            for key, columns in SCHEMAS.items():
                ws = wb.create_sheet(key)
                ws.append(columns)
                for record in data[key]:
                    ws.append([record.get(c) for c in columns])
                    for cell in ws[ws.max_row]:
                        if isinstance(cell.value, str):
                            cell.data_type = "s"  # Texto de usuario, nunca fórmulas Excel.
            wb.save(temp)
            wb.close()
            for attempt in range(5):
                try:
                    os.replace(temp, path)
                    break
                except PermissionError:
                    if attempt == 4:
                        raise
                    # Windows puede mantener un bloqueo breve durante el escaneo
                    # de un archivo recién creado. Nunca truncar el original.
                    time.sleep(0.05 * (attempt + 1))
        finally:
            wb.close()
            if temp.exists():
                try:
                    temp.unlink()
                except PermissionError:
                    pass  # No ocultar el error original de escritura.
        return load(path)


def cop(value):
    return "$ " + f"{value or 0:,.0f}".replace(",", ".")


def period(data, year, month=None):
    start = date(year, month or 1, 1).isoformat()
    end_month = month or 12
    end = date(year, end_month, monthrange(year, end_month)[1]).isoformat()
    selected = lambda rows: [r for r in rows if start <= r["fecha"] <= end]
    income, expenses, movements = (selected(data[k]) for k in ("ingresos", "gastos", "movimientos"))
    result = {}
    for label, rows in [("ingresos", income), ("gastos", expenses)]:
        for field in ["planeado", "real"]:
            result[f"{label}_{field}"] = sum(r[f"valor_{field}"] for r in rows)
        result[f"{label}_diferencia"] = (result[f"{label}_real"] - result[f"{label}_planeado"]) * (1 if label == "ingresos" else -1)
    result["balance_planeado"] = result["ingresos_planeado"] - result["gastos_planeado"]
    result["balance_real"] = result["ingresos_real"] - result["gastos_real"]
    result["balance_diferencia"] = result["balance_real"] - result["balance_planeado"]
    cash_payments = lambda rows: sum(m["valor"] for m in rows if m["tipo_movimiento"] != "Interés")
    result["abonos"] = sum(m["valor"] for m in movements if m["tipo_movimiento"] == "Abono dirigido a capital")
    result["pagos"] = cash_payments(movements) - result["abonos"]
    result["disponible"] = result["balance_real"] - cash_payments(movements)
    result["acumulada"] = (sum(r["valor_real"] for r in data["ingresos"] if r["fecha"] <= end)
                            - sum(r["valor_real"] for r in data["gastos"] if r["fecha"] <= end)
                            - cash_payments([m for m in data["movimientos"] if m["fecha"] <= end]))
    previous = (date.fromisoformat(start) - timedelta(days=1)).isoformat()
    result["saldo_inicial"] = sum(debt_balances(data, previous).values())
    result["nueva_deuda"] = sum(d["saldo_inicial"] for d in data["deudas"] if start <= d["fecha_inicio"] <= end)
    result["saldo_actual"] = sum(debt_balances(data, end).values())
    result["reduccion_deuda"] = result["saldo_inicial"] + result["nueva_deuda"] - result["saldo_actual"]
    return result


def categories(data, year, month):
    prefix = f"{year:04d}-{month:02d}"
    grouped = {}
    for r in data["gastos"]:
        if r["fecha"].startswith(prefix):
            item = grouped.setdefault(r["categoria"], {"categoria": category_label(data, r["categoria"]), "planeado": 0, "real": 0})
            item["planeado"] += r["valor_planeado"]
            item["real"] += r["valor_real"]
    for item in grouped.values():
        item["diferencia"] = item["planeado"] - item["real"]
        item["ejecucion"] = round(item["real"] / item["planeado"] * 100, 1) if item["planeado"] else None
    return sorted(grouped.values(), key=lambda r: r["real"], reverse=True)


def project(data, cutoff=None):
    cutoff = cutoff or today().isoformat()
    balance = sum(debt_balances(data, cutoff).values())
    rows = []
    for r in sorted(data["proyecciones"], key=lambda r: r["fecha"]):
        if r["fecha"][:7] <= cutoff[:7]:
            continue
        if r["abono_extraordinario"] > balance:
            raise ValueError(f"El abono proyectado de {r['fecha'][:7]} supera la deuda pendiente proyectada.")
        balance -= r["abono_extraordinario"]
        rows.append(dict(r, liquidez_restante=r["disponible_proyectado"] - r["abono_extraordinario"], saldo_proyectado=balance))
    return rows
