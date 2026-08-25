"""
Centro de Conocimiento — lectura y escritura de la pestaña CONOCIMIENTO (y su
tabla de archivos adjuntos CONOCIMIENTO_ARCHIVOS) en el mismo seguimiento.xlsx.

Mismo patrón que data.py: toda escritura abre el workbook con openpyxl, ubica
columnas por nombre de encabezado, identifica la fila por conocimiento_id
(único, nunca cambia), escribe, y guarda de forma síncrona antes de responder.
Ningún otro módulo debe tocar estas hojas directamente.
"""
from __future__ import annotations

import datetime as dt
import re
from pathlib import Path

import openpyxl
import pandas as pd

import data as data_mod

KNOWLEDGE_SHEET = "CONOCIMIENTO"
KNOWLEDGE_COLUMNS = [
    "conocimiento_id", "titulo", "descripcion_breve", "contenido", "categoria",
    "estado", "ambito", "proyectos", "etiquetas", "fuente",
    "actividades_relacionadas", "hallazgos_relacionados", "conceptos", "objetivo_estudio",
    "fecha_creacion", "fecha_actualizacion",
]
KNOWLEDGE_DATE_COLS = ["fecha_creacion", "fecha_actualizacion"]

ARCHIVOS_SHEET = "CONOCIMIENTO_ARCHIVOS"
ARCHIVOS_COLUMNS = ["conocimiento_id", "nombre_archivo", "ruta", "fecha_subida"]

CATEGORIAS = ["Estudio", "Soluciones", "Errores y aprendizajes", "Conceptos",
              "Procedimientos", "Ideas", "Referencias", "Notas"]
ESTADOS = ["En estudio", "En progreso", "Aprendido", "Aplicado",
           "Pendiente de revisar", "Archivado"]
AMBITOS = ["Global", "Proyecto"]

# Los archivos adjuntos viven en una subcarpeta del mismo disco donde ya vive
# seguimiento.xlsx (local o el disco persistente de Render) — no se introduce
# almacenamiento externo nuevo.
ARCHIVOS_DIR = data_mod.EXCEL_PATH.parent / "conocimiento_archivos"


def _ensure_archivos_dir() -> None:
    ARCHIVOS_DIR.mkdir(parents=True, exist_ok=True)


def parse_conceptos(texto: str | None) -> list[dict]:
    """Convierte el texto libre del campo 'conceptos' (p.ej.
    'Routing:100, Pydantic:80, Dependency Injection:40') en una lista de
    {nombre, progreso}. Un concepto sin ':porcentaje' se toma como 0%."""
    if texto is None or (isinstance(texto, float) and pd.isna(texto)) or not str(texto).strip():
        return []
    conceptos = []
    for parte in str(texto).split(","):
        parte = parte.strip()
        if not parte:
            continue
        nombre, _, pct_txt = parte.rpartition(":")
        if not nombre:
            nombre, pct_txt = parte, ""
        digitos = re.sub(r"[^\d]", "", pct_txt)
        pct = max(0, min(100, int(digitos))) if digitos else 0
        conceptos.append({"nombre": nombre.strip(), "progreso": pct})
    return conceptos


def progreso_promedio(texto: str | None) -> int | None:
    """Progreso 0-100 promediado entre los conceptos declarados, o None si
    no hay ninguno (para no confundir '0% real' con 'sin conceptos')."""
    conceptos = parse_conceptos(texto)
    if not conceptos:
        return None
    return round(sum(c["progreso"] for c in conceptos) / len(conceptos))


def load_knowledge(path: Path | str | None = None) -> pd.DataFrame:
    path = path if path is not None else data_mod.EXCEL_PATH
    df = pd.read_excel(path, sheet_name=KNOWLEDGE_SHEET, usecols=KNOWLEDGE_COLUMNS)
    for col in KNOWLEDGE_COLUMNS:
        if col not in KNOWLEDGE_DATE_COLS:
            df[col] = df[col].apply(lambda v: v.strip() if isinstance(v, str) else v)
    for col in KNOWLEDGE_DATE_COLS:
        df[col] = pd.to_datetime(df[col], errors="coerce")
    return df.reset_index(drop=True)


def load_knowledge_files(path: Path | str | None = None) -> pd.DataFrame:
    path = path if path is not None else data_mod.EXCEL_PATH
    df = pd.read_excel(path, sheet_name=ARCHIVOS_SHEET, usecols=ARCHIVOS_COLUMNS)
    for col in ["conocimiento_id", "nombre_archivo", "ruta"]:
        df[col] = df[col].apply(lambda v: v.strip() if isinstance(v, str) else v)
    df["fecha_subida"] = pd.to_datetime(df["fecha_subida"], errors="coerce")
    return df.reset_index(drop=True)


def search_knowledge(query: str | None = None, categoria: str | None = None,
                      proyecto: str | None = None, etiqueta: str | None = None,
                      estado: str | None = None, path: Path | str | None = None) -> pd.DataFrame:
    """Búsqueda simple por texto sobre título/descripción/contenido/etiquetas,
    más filtros exactos. Sin vectores/embeddings en esta primera versión."""
    df = load_knowledge(path)
    if df.empty:
        return df
    if query:
        q = query.strip().lower()
        mascara = (
            df["titulo"].fillna("").str.lower().str.contains(q, regex=False)
            | df["descripcion_breve"].fillna("").str.lower().str.contains(q, regex=False)
            | df["contenido"].fillna("").str.lower().str.contains(q, regex=False)
            | df["etiquetas"].fillna("").str.lower().str.contains(q, regex=False)
        )
        df = df[mascara]
    if categoria:
        df = df[df["categoria"] == categoria]
    if estado:
        df = df[df["estado"] == estado]
    if proyecto:
        df = df[df["proyectos"].fillna("").str.contains(proyecto, case=False, regex=False)]
    if etiqueta:
        df = df[df["etiquetas"].fillna("").str.contains(etiqueta, case=False, regex=False)]
    return df.reset_index(drop=True)


def _abrir_para_escritura(path, accion: str):
    try:
        wb = openpyxl.load_workbook(path)
    except PermissionError:
        return None, None, None, (f"No fue posible {accion}. Verifique que el archivo Excel "
                                    "no esté abierto o bloqueado por otro usuario.")
    except FileNotFoundError:
        return None, None, None, "No se encontró el archivo seguimiento.xlsx."

    if KNOWLEDGE_SHEET not in wb.sheetnames:
        return None, None, None, "No se encontró la pestaña CONOCIMIENTO en el archivo."

    ws = wb[KNOWLEDGE_SHEET]
    headers = [cell.value for cell in ws[1]]
    try:
        col_idx = {name: headers.index(name) + 1 for name in KNOWLEDGE_COLUMNS}
    except ValueError:
        return None, None, None, "La estructura de columnas de CONOCIMIENTO cambió y no se pudo actualizar de forma segura."

    return wb, ws, col_idx, None


def next_knowledge_id(path: Path | str | None = None) -> str:
    path = path if path is not None else data_mod.EXCEL_PATH
    ids = pd.read_excel(path, sheet_name=KNOWLEDGE_SHEET, usecols="A")["conocimiento_id"].dropna().astype(str)
    nums = [int(m.group()) for i in ids if (m := re.search(r"\d+", i))]
    return f"K{(max(nums) + 1) if nums else 1:03d}"


def add_knowledge(titulo: str, descripcion_breve: str, contenido: str, categoria: str,
                   estado: str, ambito: str, proyectos: str | None, etiquetas: str | None,
                   fuente: str | None = None, actividades_relacionadas: str | None = None,
                   hallazgos_relacionados: str | None = None, conceptos: str | None = None,
                   objetivo_estudio: str | None = None,
                   path: Path | str | None = None) -> tuple[bool, str, str | None]:
    """Agrega una entrada nueva al final de CONOCIMIENTO. Devuelve también el
    ID generado (K0xx) para poder, por ejemplo, subir un archivo adjunto en el
    mismo flujo sin recargar la página."""
    path = path if path is not None else data_mod.EXCEL_PATH
    wb, ws, col_idx, error = _abrir_para_escritura(path, "agregar el conocimiento")
    if error:
        return False, error, None

    ids_existentes = [str(ws.cell(row=r, column=col_idx["conocimiento_id"]).value or "")
                       for r in range(2, ws.max_row + 1)]
    nums = [int(m.group()) for i in ids_existentes if (m := re.search(r"\d+", i))]
    new_id = f"K{(max(nums) + 1) if nums else 1:03d}"

    ahora = dt.date.today()
    fila = ws.max_row + 1
    valores = {
        "conocimiento_id": new_id, "titulo": titulo, "descripcion_breve": descripcion_breve,
        "contenido": contenido, "categoria": categoria, "estado": estado, "ambito": ambito,
        "proyectos": proyectos or None, "etiquetas": etiquetas or None, "fuente": fuente or None,
        "actividades_relacionadas": actividades_relacionadas or None,
        "hallazgos_relacionados": hallazgos_relacionados or None,
        "conceptos": conceptos or None, "objetivo_estudio": objetivo_estudio or None,
        "fecha_creacion": ahora, "fecha_actualizacion": ahora,
    }
    for nombre, valor in valores.items():
        celda = ws.cell(row=fila, column=col_idx[nombre])
        celda.value = valor
        if nombre in KNOWLEDGE_DATE_COLS:
            celda.number_format = "DD/MM/YYYY"

    try:
        wb.save(path)
    except PermissionError:
        return False, ("No fue posible agregar el conocimiento. Verifique que el archivo Excel "
                        "no esté abierto o bloqueado por otro usuario."), None
    return True, f"Conocimiento {new_id} agregado correctamente.", new_id


def update_knowledge(conocimiento_id: str, titulo: str, descripcion_breve: str, contenido: str,
                      categoria: str, estado: str, ambito: str, proyectos: str | None,
                      etiquetas: str | None, fuente: str | None = None,
                      actividades_relacionadas: str | None = None,
                      hallazgos_relacionados: str | None = None, conceptos: str | None = None,
                      objetivo_estudio: str | None = None,
                      path: Path | str | None = None) -> tuple[bool, str]:
    path = path if path is not None else data_mod.EXCEL_PATH
    wb, ws, col_idx, error = _abrir_para_escritura(path, "actualizar el conocimiento")
    if error:
        return False, error

    filas = [r for r in range(2, ws.max_row + 1)
             if str(ws.cell(row=r, column=col_idx["conocimiento_id"]).value or "") == str(conocimiento_id)]
    if not filas:
        return False, "No se encontró esa entrada de conocimiento (los datos pudieron cambiar)."
    if len(filas) > 1:
        return False, "Existe más de una fila con el mismo ID; no es posible actualizar de forma segura."

    fila = filas[0]
    valores = {
        "titulo": titulo, "descripcion_breve": descripcion_breve, "contenido": contenido,
        "categoria": categoria, "estado": estado, "ambito": ambito,
        "proyectos": proyectos or None, "etiquetas": etiquetas or None, "fuente": fuente or None,
        "actividades_relacionadas": actividades_relacionadas or None,
        "hallazgos_relacionados": hallazgos_relacionados or None,
        "conceptos": conceptos or None, "objetivo_estudio": objetivo_estudio or None,
        "fecha_actualizacion": dt.date.today(),
    }
    for nombre, valor in valores.items():
        celda = ws.cell(row=fila, column=col_idx[nombre])
        celda.value = valor
        if nombre in KNOWLEDGE_DATE_COLS:
            celda.number_format = "DD/MM/YYYY"

    try:
        wb.save(path)
    except PermissionError:
        return False, ("No fue posible actualizar el conocimiento. Verifique que el archivo Excel "
                        "no esté abierto o bloqueado por otro usuario.")
    return True, f"Conocimiento {conocimiento_id} actualizado correctamente."


def delete_knowledge(conocimiento_id: str, path: Path | str | None = None) -> tuple[bool, str]:
    """Elimina la entrada y, en cascada, sus filas de CONOCIMIENTO_ARCHIVOS y
    los archivos físicos asociados (huérfanos no sirven de nada). No se puede
    deshacer."""
    path = path if path is not None else data_mod.EXCEL_PATH
    try:
        wb = openpyxl.load_workbook(path)
    except PermissionError:
        return False, ("No fue posible eliminar el conocimiento. Verifique que el archivo Excel "
                        "no esté abierto o bloqueado por otro usuario.")
    except FileNotFoundError:
        return False, "No se encontró el archivo seguimiento.xlsx."

    if KNOWLEDGE_SHEET not in wb.sheetnames or ARCHIVOS_SHEET not in wb.sheetnames:
        return False, "No se encontraron las pestañas de CONOCIMIENTO en el archivo."

    ws = wb[KNOWLEDGE_SHEET]
    headers = [cell.value for cell in ws[1]]
    try:
        col_idx = {name: headers.index(name) + 1 for name in KNOWLEDGE_COLUMNS}
    except ValueError:
        return False, "La estructura de columnas de CONOCIMIENTO cambió y no se pudo actualizar de forma segura."

    filas = [r for r in range(2, ws.max_row + 1)
             if str(ws.cell(row=r, column=col_idx["conocimiento_id"]).value or "") == str(conocimiento_id)]
    if not filas:
        return False, "No se encontró esa entrada de conocimiento (los datos pudieron cambiar)."
    if len(filas) > 1:
        return False, "Existe más de una fila con el mismo ID; no es posible eliminar de forma segura."

    ws.delete_rows(filas[0], 1)

    wsa = wb[ARCHIVOS_SHEET]
    headers_a = [cell.value for cell in wsa[1]]
    col_idx_a = {name: headers_a.index(name) + 1 for name in ARCHIVOS_COLUMNS}
    rutas_a_borrar = []
    filas_archivo = []
    for r in range(2, wsa.max_row + 1):
        if str(wsa.cell(row=r, column=col_idx_a["conocimiento_id"]).value or "") == str(conocimiento_id):
            filas_archivo.append(r)
            rutas_a_borrar.append(wsa.cell(row=r, column=col_idx_a["ruta"]).value)
    for r in sorted(filas_archivo, reverse=True):
        wsa.delete_rows(r, 1)

    try:
        wb.save(path)
    except PermissionError:
        return False, ("No fue posible eliminar el conocimiento. Verifique que el archivo Excel "
                        "no esté abierto o bloqueado por otro usuario.")

    for ruta in rutas_a_borrar:
        if ruta:
            try:
                Path(ruta).unlink(missing_ok=True)
            except OSError:
                pass

    return True, f"Conocimiento {conocimiento_id} eliminado correctamente."


def add_knowledge_file(conocimiento_id: str, filename: str, contenido_bytes: bytes,
                        path: Path | str | None = None) -> tuple[bool, str]:
    """Guarda el archivo físico en conocimiento_archivos/{conocimiento_id}/ y
    registra la fila en CONOCIMIENTO_ARCHIVOS."""
    path = path if path is not None else data_mod.EXCEL_PATH
    wb, ws, col_idx, error = _abrir_para_escritura(path, "agregar el archivo")
    if error:
        return False, error

    if ARCHIVOS_SHEET not in wb.sheetnames:
        return False, "No se encontró la pestaña CONOCIMIENTO_ARCHIVOS en el archivo."
    wsa = wb[ARCHIVOS_SHEET]
    headers_a = [cell.value for cell in wsa[1]]
    try:
        col_idx_a = {name: headers_a.index(name) + 1 for name in ARCHIVOS_COLUMNS}
    except ValueError:
        return False, "La estructura de columnas de CONOCIMIENTO_ARCHIVOS cambió y no se pudo actualizar de forma segura."

    _ensure_archivos_dir()
    carpeta_entrada = ARCHIVOS_DIR / conocimiento_id
    carpeta_entrada.mkdir(parents=True, exist_ok=True)
    nombre_seguro = re.sub(r"[^\w.\-]", "_", filename)
    destino = carpeta_entrada / nombre_seguro
    destino.write_bytes(contenido_bytes)

    fila = wsa.max_row + 1
    ahora = dt.date.today()
    valores_a = {"conocimiento_id": conocimiento_id, "nombre_archivo": filename,
                 "ruta": str(destino), "fecha_subida": ahora}
    for nombre, valor in valores_a.items():
        celda = wsa.cell(row=fila, column=col_idx_a[nombre])
        celda.value = valor
        if nombre == "fecha_subida":
            celda.number_format = "DD/MM/YYYY"

    try:
        wb.save(path)
    except PermissionError:
        destino.unlink(missing_ok=True)
        return False, ("No fue posible agregar el archivo. Verifique que el archivo Excel "
                        "no esté abierto o bloqueado por otro usuario.")
    return True, f"Archivo '{filename}' agregado correctamente."
