import sys
import tempfile
import unittest
import openpyxl
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "app"))
import personal as p


class PersonalTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / "personal.xlsx"

    def save(self, kind, **row):
        return p.mutate(kind, row, path=self.path)

    def income(self, **kw):
        return self.save("ingresos", **dict(dict(fecha="2026-10-01", descripcion="Salario", categoria=self.category("Salario", "Ingreso"), valor_planeado=1000, valor_real=900, estado="Parcial"), **kw))

    def expense(self, **kw):
        return self.save("gastos", **dict(dict(fecha="2026-10-02", descripcion="Comida", categoria=self.category("Comida", "Gasto"), valor_planeado=500, valor_real=300, estado="Pagado", tipo="Variable"), **kw))

    def category(self, name, kind="Gasto", parent=""):
        existing = next((c for c in p.load(self.path)["categorias"] if c["nombre"] == name and c["tipo"] == kind and c["padre"] == parent), None)
        if existing:
            return existing["id"]
        return self.save("categorias", nombre=name, tipo=kind, estado="Activa", padre=parent)["categorias"][-1]["id"]

    def test_category_rename_preserves_history(self):
        data = self.expense()
        before = data["gastos"]
        category = data["categorias"][0]
        data = self.save("categorias", **dict(category, nombre="Alimentación"))
        self.assertEqual(data["gastos"], before)
        self.assertEqual(p.categories(data, 2026, 10)[0]["categoria"], "Alimentación")
        self.assertEqual(p.period(data, 2026)["gastos_real"], 300)

    def test_daily_expenses_update_totals_without_changing_budget(self):
        data = self.expense(valor_planeado=100000, valor_real=5000)
        category = data["gastos"][0]["categoria"]
        original = data["gastos"]
        self.save("gastos_diarios", fecha="2026-10-04", descripcion="Taxi", categoria=category, valor=20000)
        data = self.save("gastos_diarios", fecha="2026-10-04", descripcion="Taxi", categoria=category, valor=20000)
        self.assertEqual(data["gastos"], original)
        self.assertEqual(p.period(data, 2026, 10)["gastos_real"], 45000)
        self.assertEqual(p.period(data, 2026, 10)["gastos_planeado"], 100000)
        self.assertEqual(p.period(data, 2026, 10)["disponible"], -45000)
        self.assertEqual(p.categories(data, 2026, 10)[0]["real"], 45000)
        self.assertEqual(p.daily_totals(data, 2026, 10)[0]["valor"], 40000)
        row = data["gastos_diarios"][0]
        data = self.save("gastos_diarios", **dict(row, fecha="2026-11-01", valor=10000))
        self.assertEqual(p.period(data, 2026, 10)["gastos_real"], 25000)
        self.assertEqual(p.period(data, 2026, 11)["gastos_real"], 10000)
        self.assertEqual(p.period(data, 2026)["gastos_real"], 35000)
        self.assertEqual(p.period(data, 2026, 11)["acumulada"], -35000)
        data = p.mutate("gastos_diarios", delete_id=row["id"], path=self.path)
        self.assertEqual(p.period(data, 2026)["gastos_real"], 25000)

    def test_daily_categories_and_invalid_values(self):
        category = self.category("Taxis")
        row = dict(fecha="2026-10-04", descripcion="Taxi", categoria=category, valor=20000)
        data = self.save("gastos_diarios", **row)
        for value in (0, -1, 1.5):
            with self.assertRaises(ValueError):
                self.save("gastos_diarios", **dict(row, valor=value))
        with self.assertRaises(ValueError):
            self.save("gastos_diarios", **dict(row, categoria=self.category("Salario", "Ingreso")))
        with self.assertRaises(ValueError):
            p.mutate("categorias", delete_id=category, path=self.path)
        cat = next(c for c in data["categorias"] if c["id"] == category)
        self.save("categorias", **dict(cat, estado="Inactiva", nombre="Transporte"))
        with self.assertRaises(ValueError):
            self.save("gastos_diarios", **row)
        data = self.save("gastos_diarios", **dict(data["gastos_diarios"][0], valor=25000))
        self.assertEqual(p.daily_totals(data, 2026, 10)[0]["categoria"], "Transporte")

    def test_existing_workbook_without_daily_sheet_preserves_values(self):
        before = self.expense()
        wb = openpyxl.load_workbook(self.path)
        del wb["gastos_diarios"]
        wb.save(self.path)
        wb.close()
        raw = self.path.read_bytes()
        data = p.load(self.path)
        self.assertEqual(data, before)
        self.assertEqual(self.path.read_bytes(), raw)
        data = self.save("gastos_diarios", fecha="2026-10-04", descripcion="Taxi",
                         categoria=before["gastos"][0]["categoria"], valor=20000)
        self.assertEqual(data["gastos"], before["gastos"])

    def test_inactive_category_preserves_existing_but_blocks_new(self):
        data = self.expense()
        category = data["categorias"][0]
        data = self.save("categorias", **dict(category, estado="Inactiva"))
        self.assertEqual(p.category_options(data, "gastos"), [])
        self.assertEqual(len(p.category_options(data, "gastos", category["id"])), 1)
        with self.assertRaisesRegex(ValueError, "activa"):
            self.expense(descripcion="Otro gasto")
        updated = self.save("gastos", **dict(data["gastos"][0], observaciones="Nota histórica"))
        self.assertEqual(updated["gastos"][0]["categoria"], category["id"])
        data = self.save("categorias", **dict(category, estado="Activa"))
        self.assertEqual(len(p.category_options(data, "gastos")), 1)

    def test_reassign_category_and_type_validation(self):
        data = self.expense()
        old = data["categorias"][0]
        new_id = self.category("Mercado")
        data = self.save("gastos", **dict(data["gastos"][0], categoria=new_id))
        self.assertEqual(data["gastos"][0]["categoria"], new_id)
        income_id = self.category("Extras", "Ingreso")
        with self.assertRaises(ValueError):
            self.save("gastos", **dict(data["gastos"][0], categoria=income_id))
        used = next(c for c in data["categorias"] if c["id"] == new_id)
        with self.assertRaises(ValueError):
            self.save("categorias", **dict(used, tipo="Ingreso"))
        p.mutate("categorias", delete_id=old["id"], path=self.path)

    def test_used_category_cannot_be_deleted(self):
        data = self.income()
        before = self.path.read_bytes()
        with self.assertRaisesRegex(ValueError, "histórico"):
            p.mutate("categorias", delete_id=data["categorias"][0]["id"], path=self.path)
        self.assertEqual(before, self.path.read_bytes())

    def test_subcategories_cycles_and_parent_deactivation(self):
        parent = self.category("Hogar")
        child = self.category("Servicios", parent=parent)
        data = self.expense(categoria=child)
        self.assertEqual(p.categories(data, 2026, 10)[0]["categoria"], "Hogar / Servicios")
        parent_row = next(c for c in data["categorias"] if c["id"] == parent)
        with self.assertRaisesRegex(ValueError, "ciclos"):
            self.save("categorias", **dict(parent_row, padre=child))
        with self.assertRaisesRegex(ValueError, "subcategorías"):
            p.mutate("categorias", delete_id=parent, path=self.path)
        data = self.save("categorias", **dict(parent_row, estado="Inactiva"))
        self.assertFalse(p.category_active(data, child))
        with self.assertRaises(ValueError):
            self.expense(categoria=child, descripcion="Nuevo")
        with self.assertRaises(ValueError):
            self.category("Salario", "Ingreso", parent=parent)

    def test_category_names_unique_within_type_and_parent(self):
        self.category("Comida")
        with self.assertRaises(ValueError):
            self.save("categorias", nombre=" comida ", tipo="Gasto", estado="Activa")
        self.category("Comida", "Ingreso")
        parent = self.category("Viajes")
        self.category("Comida", parent=parent)
        self.assertEqual(len(p.load(self.path)["categorias"]), 4)

    def test_legacy_category_migration_is_stable_and_preserves_values(self):
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for key, columns in p.SCHEMAS.items():
            if key == "categorias":
                continue
            ws = wb.create_sheet(key)
            ws.append(columns)
            if key in ("ingresos", "gastos"):
                row = dict(id=key, fecha="2026-10-01", descripcion="Anterior", categoria="Otros", valor_planeado=100, valor_real=80,
                           estado="Planeado", tipo="Variable", observaciones="")
                ws.append([row[c] for c in columns])
        wb.save(self.path)
        wb.close()
        before = self.path.read_bytes()
        data = p.load(self.path)
        self.assertEqual(p.revision(data), p.revision(p.load(self.path)))
        self.assertEqual(self.path.read_bytes(), before)
        self.assertEqual(len(data["categorias"]), 2)
        self.assertNotEqual(data["ingresos"][0]["categoria"], data["gastos"][0]["categoria"])
        category = data["categorias"][0]
        updated = self.save("categorias", **dict(category, nombre="Renombrada"))
        self.assertEqual(updated["gastos"], data["gastos"])
        self.assertEqual(updated["ingresos"], data["ingresos"])

    def debt(self, **kw):
        data = self.save("deudas", **dict(dict(nombre="Compra B", entidad="Davivienda", saldo_inicial=5000000, tasa_EA=29.23,
            numero_cuotas=24, cuotas_restantes=24, fecha_inicio="2026-10-01", estado="Activa"), **kw))
        return data["deudas"][-1]["id"]

    def movement(self, debt, **kw):
        return self.save("movimientos", **dict(dict(fecha="2026-10-10", tipo_movimiento="Abono dirigido a capital", valor=1000000, deuda=debt), **kw))

    def test_income_totals(self):
        self.income()
        self.income(descripcion="Arriendo", valor_real=100, valor_planeado=100)
        totals = p.period(p.load(self.path), 2026, 10)
        self.assertEqual(totals["ingresos_real"], 1000)
        self.assertEqual(totals["ingresos_planeado"], 1100)

    def test_expense_totals(self):
        self.expense()
        self.expense(descripcion="Mercado", valor_real=100)
        self.assertEqual(p.period(p.load(self.path), 2026, 10)["gastos_real"], 400)

    def test_balance_and_differences(self):
        self.income()
        self.expense()
        s = p.period(p.load(self.path), 2026, 10)
        self.assertEqual((s["balance_real"], s["balance_planeado"], s["balance_diferencia"]), (600, 500, 100))
        self.assertEqual((s["ingresos_diferencia"], s["gastos_diferencia"]), (-100, 200))
        self.assertEqual(p.categories(p.load(self.path), 2026, 10)[0]["diferencia"], 200)

    def test_capital_payment_and_persistence(self):
        debt = self.debt()
        data = self.movement(debt)
        self.assertEqual(data["movimientos"][0]["capital"], 1000000)
        self.assertEqual(p.load(self.path)["deudas"][0]["saldo_actual"], 4000000)

    def test_negative_balance_rejected_atomically(self):
        debt = self.debt()
        before = self.path.read_bytes()
        with self.assertRaisesRegex(ValueError, "supera"):
            self.movement(debt, valor=5000001)
        self.assertEqual(before, self.path.read_bytes())

    def test_interest_and_normal_payment_are_separate(self):
        debt = self.debt()
        self.movement(debt, tipo_movimiento="Interés", valor=50000)
        self.movement(debt, tipo_movimiento="Cuota normal", valor=200000, capital=150000)
        data = p.load(self.path)
        self.assertEqual(data["deudas"][0]["saldo_actual"], 4900000)
        self.assertEqual(p.period(data, 2026, 10)["disponible"], -200000)

    def test_edit_and_delete_recalculate(self):
        debt = self.debt()
        data = self.movement(debt)
        m = data["movimientos"][0]
        data = self.save("movimientos", **dict(m, valor=2000000))
        self.assertEqual(data["deudas"][0]["saldo_actual"], 3000000)
        data = p.mutate("movimientos", delete_id=m["id"], path=self.path)
        self.assertEqual(data["deudas"][0]["saldo_actual"], 5000000)

    def test_monthly_and_annual_aggregation(self):
        self.income()
        self.income(fecha="2026-11-01", valor_real=100)
        self.income(fecha="2027-01-01", valor_real=200)
        data = p.load(self.path)
        self.assertEqual(p.period(data, 2026, 10)["ingresos_real"], 900)
        self.assertEqual(p.period(data, 2026)["ingresos_real"], 1000)
        self.assertEqual(p.period(data, 2027)["ingresos_real"], 200)

    def test_historical_debt_balance(self):
        debt = self.debt()
        self.movement(debt, fecha="2026-11-01")
        data = p.load(self.path)
        self.assertEqual(p.period(data, 2026, 10)["saldo_actual"], 5000000)
        self.assertEqual(p.period(data, 2026, 11)["saldo_actual"], 4000000)
        self.assertEqual(p.period(data, 2026)["reduccion_deuda"], 1000000)

    def test_projection_liquidity_and_debt(self):
        self.debt()
        data = p.load(self.path)
        data["proyecciones"] = [dict(id="p", fecha="2026-11-01", disponible_proyectado=2700000, abono_extraordinario=1000000)]
        row = p.project(data, "2026-10-04")[0]
        self.assertEqual(row["liquidez_restante"], 1700000)
        self.assertEqual(row["saldo_proyectado"], 4000000)
        data["proyecciones"][0]["abono_extraordinario"] = 6000000
        with self.assertRaises(ValueError):
            p.project(data, "2026-10-04")

    def test_invalid_values_dates_and_nulls(self):
        for value in (-1, float("nan"), float("inf"), 1.5, "abc"):
            with self.subTest(value=value), self.assertRaises(ValueError):
                self.income(valor_real=value)
        with self.assertRaises(ValueError):
            self.income(fecha="2026-02-30")
        data = self.income(valor_real=None)
        self.assertEqual(data["ingresos"][0]["valor_real"], 0)
        self.assertEqual(p.cop(12500000), "$ 12.500.000")

    def test_duplicate_and_stale_write(self):
        previous = p.revision(p.empty())
        data = self.income()
        with self.assertRaisesRegex(ValueError, "idéntico"):
            self.income()
        with self.assertRaisesRegex(ValueError, "otra pestaña"):
            p.mutate("ingresos", data["ingresos"][0], expected=previous, path=self.path)

    def test_debt_delete_requires_removing_movements(self):
        debt = self.debt()
        self.movement(debt)
        with self.assertRaisesRegex(ValueError, "primero"):
            p.mutate("deudas", delete_id=debt, path=self.path)

    def test_total_payment_must_match_and_backdated_change_validated(self):
        debt = self.debt()
        with self.assertRaisesRegex(ValueError, "coincidir"):
            self.movement(debt, tipo_movimiento="Pago total")
        self.movement(debt, tipo_movimiento="Pago total", valor=5000000)
        with self.assertRaises(ValueError):
            self.movement(debt, fecha="2026-10-05", valor=100)

    def test_zero_budget_and_cash_not_double_counted(self):
        self.income()
        self.expense(valor_planeado=0)
        debt = self.debt()
        self.movement(debt, valor=100)
        data = p.load(self.path)
        self.assertIsNone(p.categories(data, 2026, 10)[0]["ejecucion"])
        self.assertEqual(p.period(data, 2026, 10)["disponible"], 500)


if __name__ == "__main__":
    unittest.main()
